from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

SRC = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
NAVY = RGBColor(0x1E, 0x27, 0x61)
BODY = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xE5, 0xB8, 0x22)
p = Presentation(SRC)


def set_paras(shape, paras):
    tf = shape.text_frame
    tf.clear()
    for pi, runs in enumerate(paras):
        para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        for (txt, sz, bold, italic, color) in runs:
            r = para.add_run()
            r.text = txt
            f = r.font
            f.name = "Calibri"
            f.size = Pt(sz)
            f.bold = bold
            f.italic = italic
            f.color.rgb = color


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    print("MISSING", name)
    return None


# S5 After: reportedly + forces-an-answer kept as design claim
s5 = p.slides[4]
set_paras(by_name(s5, "After 103"), [
    [("After", 18, True, False, GOLD)],
    [("Pressure at Gate B opens overflow at Gate A upstream. Kumbh's alarms reportedly went "
      "unactioned; ours forces an answer: owner, timer, log.", 15, True, False, WHITE)],
])

# S2 timeline third event + caption: outcome-based only
s2 = p.slides[1]
evs = [sh for sh in s2.shapes if sh.name == "P2 ev"]
evs.sort(key=lambda s: s.top)
sh = evs[2]
tf = sh.text_frame
tf.clear()
para = tf.paragraphs[0]
for (txt, sz, bold, color) in [("After", 16, True, NAVY),
                               ("\n30 dead officially, despite live AI alerts. No action followed in time.", 14, False, BODY)]:
    r = para.add_run()
    r.text = txt
    f = r.font
    f.name = "Calibri"
    f.size = Pt(sz)
    f.bold = bold
    f.color.rgb = color
print("s2 event softened")
set_paras(by_name(s2, "TextBox 113"),
          [[("Alerts existed. Action did not follow in time.", 16, False, True, GREY)]])

p.save(SRC)
print("saved")
