from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
NAVY = RGBColor(0x1E, 0x27, 0x61)
BODY = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xE5, 0xB8, 0x22)
p = Presentation(SRC)


def set_box(shape, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for pi, runs in enumerate(paras):
        para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        para.alignment = align
        for run in runs:
            txt, sz, bold, italic, color = run[:5]
            url = run[5] if len(run) > 5 else None
            r = para.add_run()
            r.text = txt
            f = r.font
            f.name = "Calibri"
            f.size = Pt(sz)
            f.bold = bold
            f.italic = italic
            f.color.rgb = color
            if url:
                try:
                    r.hyperlink.address = url
                except Exception as e:
                    print("link fail", url, e)


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    print("MISSING", name)
    return None


# ---------- S3: prior-art framing ----------
s3 = p.slides[2]
set_box(by_name(s3, "Card 115"), [
    [("Heatmaps beat boxes", 18, True, False, NAVY)],
    [("Published research shows box detectors collapse under overlap: F1 0.07 vs 0.71 point-based. "
      "Density maps cut counting error 265 to 105 MAE. (NWPU-Crowd)", 13, False, False, BODY)],
])
print("s3 card framed as prior art")

# ---------- S5: escalation tiers + differentiator + proof ----------
s5 = p.slides[4]
set_box(by_name(s5, "TextBox 109"), [
    [("Yellow watch", 20, True, False, WHITE)],
    [("Auto-pages the zone guard with a 2-minute ack timer. Overflow gates get ready.", 15, False, False, WHITE)],
])
set_box(by_name(s5, "TextBox 111"), [
    [("Red alert", 20, True, False, WHITE)],
    [("Guards confirm on handsets; unanswered reds climb to venue command, then city control.", 15, False, False, WHITE)],
])
set_box(by_name(s5, "After 103"), [
    [("After", 18, True, False, GOLD)],
    [("Pressure at Gate B opens overflow at Gate A upstream. Kumbh's system watched without forcing an answer. "
      "Ours requires one.", 15, True, False, WHITE)],
])
set_box(by_name(s5, "TextBox 115"), [[(
    "Pilot targets: warn before critical, page guards in seconds, publish the false-alarm rate, log every step. "
    "Still to prove: thresholds across venue types, calibration without surveys.",
    14, True, False, NAVY)]])
print("s5 escalation done")

# ---------- S4: unit-clean pressure caption ----------
s4 = p.slides[3]
set_box(by_name(s4, "P4 pressurec"), [
    [("Pressure (density x movement variance) crossed its warning line before turbulence began. "
      "The accident came after. Hajj 2006 data (Johansson et al.).", 11.5, False, False, BODY)],
])
print("s4 caption cleaned")

# ---------- S6: honest status, two-line strip ----------
s6 = p.slides[5]
strip = by_name(s6, "R6 build")
strip.top = Inches(6.55)
strip.height = Inches(0.62)
set_box(strip, [
    [("Status: idea stage. Numbers above are published research we build on, not our measurements.",
      10.5, True, False, NAVY)],
    [("Build from:  ", 10, False, False, NAVY),
     ("DM-Count code (MIT)  ", 10, False, False, GREY, "https://github.com/cvlab-stonybrook/DM-Count"),
     ("CSRNet model  ", 10, False, False, GREY, "https://github.com/leeyeehoo/CSRNet-pytorch"),
     ("homography calibrator  ", 10, False, False, GREY, "https://github.com/JuanIsernGhosn/homography-calibrator"),
     ("DeepStream docs", 10, False, False, GREY, "https://docs.nvidia.com/metropolis/deepstream/7.1/text/DS_using_custom_model.html")],
])
print("s6 status done")

p.save(SRC)
print("saved")
