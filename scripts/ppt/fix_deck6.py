# Judge pass 4: density/load-binned FPS chart (real bench), explicit risk formula,
# venue graph root-cause, HITL red alert, drop FTLE. Pairs with scripts/bench_density.py.
import json
from pptx import Presentation
from pptx.chart.data import CategoryChartData

DECK = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
BENCH = r"D:\Q_project\gallery\bench_density.json"
p = Presentation(DECK)
ok, fail = [], []


def swap(tf, old, new):
    for para in tf.paragraphs:
        t = "".join(r.text for r in para.runs)
        if t.strip() == old.strip():
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            return True
    return False


def edit_shape(slide, prefix, pairs):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.text_frame.text.startswith(prefix):
            for old, new in pairs:
                if swap(sh.text_frame, old, new):
                    ok.append(old[:55])
                else:
                    fail.append(prefix + " | " + old[:55])
            return
    fail.append("SHAPE NOT FOUND: " + prefix)


# ---- 1. Chart: crowd-load buckets from real bench (inference FPS) ----
bench = json.load(open(BENCH))
fps_by_bucket = {b["bucket"]: b["fps"] for b in bench["buckets"] if b["fps"]}
order = ["0-4", "5-8", "9-12", "13+"]
vals = [fps_by_bucket.get(k) for k in order]
s4 = p.slides[3]
if any(v is None for v in vals):
    fail.append(f"bench missing buckets: {dict(zip(order, vals))}")
else:
    chart_found = False
    for sh in s4.shapes:
        if not getattr(sh, "has_chart", False):
            continue
        ch = sh.chart
        cd = CategoryChartData()
        cd.categories = [
            "Sparse\n0-4 people",
            "Light\n5-8",
            "Busy\n9-12",
            "Packed\n13+",
        ]
        cd.add_series("Checks per second", tuple(float(v) for v in vals))
        ch.replace_data(cd)
        if ch.has_title:
            tf = ch.chart_title.text_frame
            runs = tf.paragraphs[0].runs
            old_size = runs[0].font.size if runs else None
            tf.text = "Checks per second as the crowd grows"
            if old_size and tf.paragraphs[0].runs:
                tf.paragraphs[0].runs[0].font.size = old_size
        vmax = max(float(v) for v in vals)
        ch.value_axis.maximum_scale = float(max(55.0, (int(vmax / 10) + 2) * 10))
        ok.append(f"chart FPS by load {vals}")
        chart_found = True
        break
    if not chart_found:
        fail.append("CHART NOT FOUND on slide 4")

edit_shape(s4, "Tested on real crowd video", [
    ("Tested on real crowd video from three places, all on one gaming laptop. Higher means smoother.",
     "Model and tracker time only, same laptop. FPS stays above 30 even when the frame fills with people."),
])

# ---- Slide 3: venue graph + explicit risk formula ----
s3 = p.slides[2]
edit_shape(s3, "Works with normal cameras", [
    ("Any ordinary camera works. High, wide angles give the clearest density per square metre.",
     "Cameras link into one venue graph, so a jam is traced back to where it started. High, wide angles measure density best."),
])
edit_shape(s3, "Crowd metrics, not faces", [
    ("Core computes density, flow and stillness per zone. No identity is needed for a crush warning.",
     "Risk = density x stillness x counterflow. Stillness only escalates risk when the zone is also packed."),
])
edit_shape(s3, "In plain words", [
    ("In plain words: the camera watches, the software flags where the crowd stops moving, and it warns guards before a crush builds.",
     "In plain words: cameras share one venue map, the software flags where the crowd stops moving, and it warns guards before a crush builds."),
])

# ---- Slide 5: HITL red alert + root-cause in After card ----
s5 = p.slides[4]
edit_shape(s5, "Red alert", [
    ("Exits open, crowd diverted, announcement made.",
     "Guards get steps on their handsets; a human confirms before any PA or door release."),
])
edit_shape(s5, "After", [
    ("One screen ranks every zone by risk: security moves from reactive to proactive.",
     "One screen ranks zones and traces jams to their origin: security moves from reactive to proactive."),
])

# ---- Slide 6: drop FTLE, claim optical-flow honestly ----
s6 = p.slides[5]
edit_shape(s6, "Predicting crushes", [
    ("Crowd science models (AURORA and FTLE) spot danger early. A Delhi trial cut response time by 37 percent.",
     "The AURORA warning model plus our live optical-flow stall meter spot danger early. A Delhi trial cut response time by 37 percent."),
])
edit_shape(s6, "Ideas:", [
    ("Ideas: AURORA warning model, FTLE bottlenecks, STAR-Crowd Delhi trial",
     "Ideas: AURORA warning model, optical-flow stall metrics, STAR-Crowd Delhi trial"),
])

p.save(DECK)
print("OK:", len(ok))
for x in ok:
    print("  +", x)
print("FAIL:", len(fail))
for x in fail:
    print("  -", x)
