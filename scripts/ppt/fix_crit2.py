from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
NAVY = RGBColor(0x1E, 0x27, 0x61)
BODY = RGBColor(0x11, 0x11, 0x11)
p = Presentation(SRC)


def set_box(shape, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for pi, runs in enumerate(paras):
        para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        para.alignment = align
        for (txt, sz, bold, italic, color) in runs:
            r = para.add_run()
            r.text = txt
            f = r.font
            f.name = "Calibri"
            f.size = Pt(sz)
            f.bold = bold
            f.italic = italic
            f.color.rgb = color


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    print("MISSING", name)
    return None


# ---- 1. S3: side-track box becomes warns-before-stillness ----
s3 = p.slides[2]
set_box(by_name(s3, "Card 117"), [
    [("Warns before stillness", 18, True, False, NAVY)],
    [("Stop-and-go waves and rising pressure fire the alert while feet still move. "
      "Stillness is the accident, never the trigger.", 13.5, False, False, BODY)],
])
print("s3 card swapped")

# ---- 2a. S4: ring buffer + PTZ fallback ----
s4 = p.slides[3]
set_box(by_name(s4, "Card 103"), [
    [("Video stays put", 20, True, False, NAVY)],
    [("Each edge box keeps a 10-minute rolling buffer and sends small JSON numbers, never raw video. "
      "Only a confirmed Red alert exports its clip, signed, to the evidence pack.", 14, False, False, BODY)],
])
set_box(by_name(s4, "Card 104"), [
    [("Calibrated where it counts", 20, True, False, NAVY)],
    [("Fixed gates are measured once. Kumbh ran 3-month camera surveys for its crowd alarm. "
      "PTZ or moved cameras fall back to relative levels until re-measured.", 14, False, False, BODY)],
])
print("s4 cards swapped")

# ---- 2b. S5: police evidence pack ----
s5 = p.slides[4]
for sh in s5.shapes:
    if sh.name == "Card 105" and sh.has_text_frame and "Police" in sh.text_frame.text:
        set_box(sh, [
            [("Police", 20, True, False, NAVY)],
            [("Every alert, zone and action exports as one evidence pack in a click, "
              "with the signed 10-minute incident clip.", 15, False, False, BODY)],
        ])
        print("s5 police swapped")

# ---- scan: any leftover dossier/box-ID/face mentions? ----
for i, sl in enumerate(p.slides, 1):
    for sh in sl.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text
            for kw in ["dossier", "box ID", "Side track", "Face file", "face-search", "insightface", "faiss"]:
                if kw.lower() in t.lower():
                    print(f"LEFTOVER s{i} {sh.name}: {kw} :: {t[:60]!r}")

p.save(SRC)
print("saved")
