from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
NAVY = RGBColor(0x1E, 0x27, 0x61)
BODY = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
p = Presentation(SRC)


def set_box(shape, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for pi, runs in enumerate(paras):
        para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        para.alignment = align
        for run in runs:
            txt, sz, bold, italic, color = run[:5]
            url = run[5] if len(run) > 5 else None
            r = para.add_run()
            r.text = txt
            f = r.font
            f.name = "Calibri"
            f.size = Pt(sz)
            f.bold = bold
            f.italic = italic
            f.color.rgb = color
            if url:
                try:
                    r.hyperlink.address = url
                except Exception as e:
                    print("link fail", url, e)


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    print("MISSING", name)
    return None


# 1. S6 Kumbh+Hathras card: add source links
s6 = p.slides[5]
set_box(by_name(s6, "Card 108"), [
    [("Kumbh and Hathras", 16, True, False, WHITE)],
    [("30 dead, Jan 2025, with AI live. 121 dead, Jul 2024, 80k permit vs 250k crowd.", 13, False, False, WHITE)],
    [("Express on Kumbh tech  ", 10.5, False, False, WHITE,
      "https://indianexpress.com/article/long-reads/maha-kumbh-goes-digital-2700-cctvs-ai-to-track-crowd-underwater-drones-9770545/"),
     ("AP on Hathras", 10.5, False, False, WHITE,
      "https://apnews.com/article/india-stampede-deaths-religious-gathering-5ef739c5cbe8bb8b732c0b3099173810")],
])
print("s6 sources added")

# 2. S4 DPDP soften
s4 = p.slides[3]
set_box(by_name(s4, "Card 105"), [
    [("Private by design", 20, True, False, NAVY)],
    [("Heatmaps cannot show a face. Saved video is blurred, wiped after 7 days, zero identity in core. "
      "Police clips are access-logged, designed to align with DPDP Act 2023 principles.", 14, False, False, BODY)],
])
# 3. pilot-covers -> pilot-will-test
set_box(by_name(s4, "P4 edgef"), [
    [("One Jetson-class box (about $250, 15 watts) per 1 to 4 cameras. Reuses existing CCTV. "
      "Pilot will test night, rain, oblique angles.", 13, False, False, GREY)],
])
print("s4 softened")

# 5. S5 pilot clarity: delivery vs ack window
s5 = p.slides[4]
set_box(by_name(s5, "TextBox 115"), [[(
    "Pilot targets: warn before critical, alerts reach guards in seconds, publish the false-alarm rate, "
    "log every step. Still to prove: thresholds across venue types, calibration without surveys.",
    14, True, False, NAVY)]])
print("s5 clarified")

p.save(SRC)
print("saved")
