from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

SRC = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
NAVY = RGBColor(0x1E, 0x27, 0x61)
BODY = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
p = Presentation(SRC)
s2 = p.slides[1]


def set_text(name, paras):
    for sh in s2.shapes:
        if sh.name == name and sh.has_text_frame:
            tf = sh.text_frame
            tf.clear()
            for pi, runs in enumerate(paras):
                para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                for (txt, sz, bold, italic, color) in runs:
                    r = para.add_run()
                    r.text = txt
                    f = r.font
                    f.name = "Calibri"
                    f.size = Pt(sz)
                    f.bold = bold
                    f.italic = italic
                    f.color.rgb = color
            print("set", name)
            return
    print("MISSING", name)


set_text("P2 timet", [[("Kumbh, 29 January 2025", 20, True, False, NAVY)]])
set_text("Stat 102", [[("Hathras 2024", 20, True, False, NAVY)],
                      [("121 dead. 80k permit, 250k crowd. Exits blocked.", 15, False, False, BODY)]])
set_text("Stat 103", [[("Mumbai 2017", 20, True, False, NAVY)],
                      [("23 dead on a station footbridge. Rain, rush hour, no metering.", 15, False, False, BODY)]])
set_text("TextBox 113", [[("The system saw it coming. No gate opened.", 16, False, True, GREY)]])

# timeline events: find P2 ev boxes in top-to-bottom order, rewrite
evs = [sh for sh in s2.shapes if sh.name == "P2 ev"]
evs.sort(key=lambda s: s.top)
new = [
    [[("Months before", 16, True, False, NAVY), ("2,751 cameras and AI alerts go live.", 14, False, False, BODY)]],
    [[("01:00", 16, True, False, NAVY), ("Barricades jumped at the Sangam nose. Crores press forward.", 14, False, False, BODY)]],
    [[("After", 16, True, False, NAVY), ("30 dead officially. The alarm watched. Nobody acted.", 14, False, False, BODY)]],
]
for sh, paras in zip(evs, new):
    tf = sh.text_frame
    tf.clear()
    for pi, runs in enumerate(paras):
        para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        for (txt, sz, bold, italic, color) in runs:
            r = para.add_run()
            txt2 = ("" if txt in ("Months before", "01:00", "After") else "\n") + txt
            r.text = txt2
            f = r.font
            f.name = "Calibri"
            f.size = Pt(sz)
            f.bold = bold
            f.italic = italic
            f.color.rgb = color
    print("event rewritten")

p.save(SRC)
print("saved")
