from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
NAVY = RGBColor(0x1E, 0x27, 0x61)
BODY = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
p = Presentation(SRC)


def set_box(shape, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for pi, runs in enumerate(paras):
        para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        para.alignment = align
        for run in runs:
            txt, sz, bold, italic, color = run[:5]
            url = run[5] if len(run) > 5 else None
            r = para.add_run()
            r.text = txt
            f = r.font
            f.name = "Calibri"
            f.size = Pt(sz)
            f.bold = bold
            f.italic = italic
            f.color.rgb = color
            if url:
                try:
                    r.hyperlink.address = url
                except Exception as e:
                    print("link fail", url, e)


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    print("MISSING", name)
    return None


# ---------- S3 ----------
s3 = p.slides[2]
set_box(by_name(s3, "Card 115"), [
    [("Heatmaps beat boxes", 18, True, False, NAVY)],
    [("Box detectors collapse under overlap: F1 0.07 vs 0.71 point-based. "
      "Density maps cut counting error 265 to 105 MAE. (NWPU-Crowd)", 13, False, False, BODY)],
])
set_box(by_name(s3, "Card 117"), [
    [("Warns before stillness", 18, True, False, NAVY)],
    [("Stop-and-go waves and rising pressure fire the alert while feet still move. "
      "Shown on Hajj 2006 footage; multi-site trials are the pilot's job.", 13, False, False, BODY)],
])
set_box(by_name(s3, "P3 densityc"), [
    [("Head-density surface: each peak is one head, the sum is the count. (Johansson 2008).",
      10.5, False, False, GREY)],
])
print("s3 done")

# ---------- S4 ----------
s4 = p.slides[3]
set_box(by_name(s4, "Card 104"), [
    [("Calibrated where it counts", 20, True, False, NAVY)],
    [("Fixed gates are measured once; Kumbh ran 3-month surveys for its alarm. "
      "Open grounds and PTZ cameras run relative levels plus entry counts.", 14, False, False, BODY)],
])
set_box(by_name(s4, "Card 105"), [
    [("Private by design", 20, True, False, NAVY)],
    [("Heatmaps cannot show a face. Saved video is blurred, wiped after 7 days, zero identity in core. "
      "Police clips are access-logged under DPDP Act 2023 rules.", 14, False, False, BODY)],
])
set_box(by_name(s4, "P4 pressurec"), [
    [("Pressure p = density x velocity variance. Hajj 2006: turbulence onset at 0.02 per s2, "
      "accident after. (Johansson et al.)", 11.5, False, False, BODY)],
])
set_box(by_name(s4, "P4 edgef"), [
    [("One Jetson-class box (about $250, 15 watts) per 1 to 4 cameras. Reuses existing CCTV. "
      "Pilot covers night, rain, oblique angles.", 13, False, False, GREY)],
])
print("s4 done")

# ---------- S5: escalation teeth ----------
s5 = p.slides[4]
set_box(by_name(s5, "TextBox 109"), [
    [("Yellow watch", 20, True, False, WHITE)],
    [("Auto-pages the zone guard with an ack timer. Overflow gates get ready.", 15, False, False, WHITE)],
])
set_box(by_name(s5, "TextBox 111"), [
    [("Red alert", 20, True, False, WHITE)],
    [("Guards get steps on their handsets and confirm. Unanswered reds climb the chain automatically.", 15, False, False, WHITE)],
])
set_box(by_name(s5, "After 103"), [
    [("After", 18, True, False, RGBColor(0xE5, 0xB8, 0x22))],
    [("Pressure at Gate B opens overflow at Gate A upstream. Every red has an owner, a timer, and a log.",
      16, True, False, WHITE)],
])
set_box(by_name(s5, "TextBox 115"), [[(
    "Pilot targets: warn before critical, page guards in seconds, publish the false-alarm rate, log every step.",
    15, True, False, NAVY)]])
print("s5 done")

# ---------- S6: build status ----------
s6 = p.slides[5]
set_box(by_name(s6, "R6 build"), [[
    ("Status: idea stage, side-track demo built. Build from:  ", 10.5, True, False, NAVY),
    ("DM-Count code (MIT)  ", 10.5, False, False, GREY, "https://github.com/cvlab-stonybrook/DM-Count"),
    ("CSRNet model  ", 10.5, False, False, GREY, "https://github.com/leeyeehoo/CSRNet-pytorch"),
    ("homography calibrator  ", 10.5, False, False, GREY, "https://github.com/JuanIsernGhosn/homography-calibrator"),
    ("DeepStream custom-model docs", 10.5, False, False, GREY, "https://docs.nvidia.com/metropolis/deepstream/7.1/text/DS_using_custom_model.html"),
]])
print("s6 done")

p.save(SRC)
print("saved")
