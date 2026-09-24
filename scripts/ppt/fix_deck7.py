"""fix_deck7: remove prototype photos (wall screenshot, tracking GIF, crowd photo)
and replace with native-shape concept diagrams for real crowd control."""
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
shutil.copy(SRC, r"C:\Users\shoun\AppData\Local\Temp\opencode\simp-backup7.pptx")

NAVY = RGBColor(0x1E, 0x27, 0x61)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
BODY = RGBColor(0x11, 0x11, 0x11)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD_FB = RGBColor(0xEB, 0xEF, 0xF9)
GREEN_FB = RGBColor(0x2F, 0xA8, 0x5A)
AMBER_FB = RGBColor(0xE8, 0xA8, 0x3D)
RED_FB = RGBColor(0xE1, 0x4B, 0x44)
A_REL = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

p = Presentation(SRC)

# --- sample actual legend-dot + card colors from slide 2 ---
s2 = p.slides[1]
sampled = {}
for sh in s2.shapes:
    if sh.name in ("Stat 102", "Dot 116", "Dot 118", "Dot 120"):
        try:
            sampled[sh.name] = sh.fill.fore_color.rgb
        except Exception as e:
            sampled[sh.name] = None
print("sampled:", {k: (str(v) if v else None) for k, v in sampled.items()})
CARD = sampled.get("Stat 102") or CARD_FB
CALM = sampled.get("Dot 116") or GREEN_FB
WATCHC = sampled.get("Dot 118") or AMBER_FB
CRIT = sampled.get("Dot 120") or RED_FB


def kill(slide, name):
    for sh in list(slide.shapes):
        if sh.name == name:
            el = sh._element
            rId = None
            blip = el.find(".//%sblip" % A_REL)
            if blip is not None:
                rId = blip.get(R_REL + "embed")
            el.getparent().remove(el)
            if rId:
                try:
                    slide.part.drop_rel(rId)
                except Exception as e:
                    print("drop_rel fail", name, rId, e)
            print("killed", name, rId)
            return True
    print("NOT FOUND", name)
    return False


def card(slide, name, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    try:
        s.adjustments[0] = 0.07
    except Exception:
        pass
    s.text_frame.clear()
    return s


def textbox(slide, name, l, t, w, h, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT):
    s = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    s.name = name
    tf = s.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    for pi, runs in enumerate(paras):
        para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        para.alignment = align
        for (txt, sz, bold, italic, color) in runs:
            r = para.add_run()
            r.text = txt
            f = r.font
            f.name = "Calibri"
            f.size = Pt(sz)
            f.bold = bold
            f.italic = italic
            f.color.rgb = color
    return s


def swap_caption(slide, old, new):
    for sh in slide.shapes:
        if sh.has_text_frame and old in sh.text_frame.text:
            for para in sh.text_frame.paragraphs:
                if old in para.text:
                    para.runs[0].text = new
                    for r in para.runs[1:]:
                        r.text = ""
            print("caption swapped in", sh.name)
            return True
    print("CAPTION NOT FOUND")
    return False


# ================= SLIDE 2: venue map =================
kill(s2, "Live wall 112")
card(s2, "VM card", 9.80, 2.85, 9.40, 5.00, CARD)
textbox(s2, "VM title", 10.10, 3.00, 8.80, 0.55, [
    [("Venue map", 20, True, False, NAVY), ("  \u2014 every zone ranked by risk", 16, False, False, NAVY)],
])
tiles = [
    ("Stage", "Calm", "all is well", CALM, "3", 10.10),
    ("Concourse", "Watch", "getting crowded", WATCHC, "2", 13.14),
    ("Gate", "Critical", "packed and stuck", CRIT, "1", 16.18),
]
for (zn, st, sub, fill, rank, x) in tiles:
    t = card(s2, "VM tile " + zn, x, 3.80, 2.45, 2.20, fill)
    tf = t.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    pa = tf.paragraphs[0]
    pa.alignment = PP_ALIGN.CENTER
    for (txt, sz, b) in [(zn, 17, True), ("%s \u2014 %s" % (st, sub), 12, False)]:
        r = pa.add_run() if pa.runs and pa.text else pa.add_run()
        r.text = ("" if txt == zn else "\n") + txt
        f = r.font
        f.name = "Calibri"
        f.size = Pt(sz)
        f.bold = b
        f.color.rgb = WHITE
    bdg = s2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.08), Inches(3.88), Inches(0.5), Inches(0.5))
    bdg.name = "VM rank " + zn
    bdg.fill.solid()
    bdg.fill.fore_color.rgb = NAVY
    bdg.line.fill.background()
    textbox(s2, "VM rankt " + zn, x + 0.08, 3.88, 0.5, 0.5, [[(rank, 15, True, False, WHITE)]],
            anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
for ax in (12.62, 15.66):
    a = s2.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(ax), Inches(4.65), Inches(0.45), Inches(0.5))
    a.name = "VM arrow"
    a.fill.solid()
    a.fill.fore_color.rgb = NAVY
    a.line.fill.background()
textbox(s2, "VM foot", 10.10, 6.15, 8.80, 1.50, [
    [("Arrows follow the crowd. Risk ranks the other way: Gate first.", 15, False, False, GREY)],
])
swap_caption(s2, "The live wall, running on real station footage.",
             "Guard screen: one map with every zone ranked by risk.")

# ================= SLIDE 3: venue graph =================
s3 = p.slides[2]
kill(s3, "Tracking 114")
card(s3, "VG card", 0.80, 6.55, 7.40, 4.00, CARD)
textbox(s3, "VG title", 1.05, 6.70, 6.90, 0.50, [
    [("One venue graph", 18, True, False, NAVY)],
])
cams = [("Gate cam", 8.05), ("Stage cam", 8.90), ("Exit cam", 9.75)]
for i, (lbl, cy) in enumerate(cams, 1):
    c = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.15), Inches(cy - 0.35), Inches(0.70), Inches(0.70))
    c.name = "VG cam %d" % i
    c.fill.solid()
    c.fill.fore_color.rgb = NAVY
    c.line.fill.background()
    textbox(s3, "VG camt %d" % i, 1.15, cy - 0.35, 0.70, 0.70, [[(str(i), 16, True, False, WHITE)]],
            anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    textbox(s3, "VG caml %d" % i, 2.00, cy - 0.30, 1.30, 0.60, [[(lbl, 13, False, False, NAVY)]],
            anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT)
    bar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.35), Inches(cy - 0.02), Inches(0.75), Inches(0.04))
    bar.name = "VG link %d" % i
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
hub = card(s3, "VG hub", 4.10, 8.45, 1.55, 1.05, NAVY)
textbox(s3, "VG hubt", 4.10, 8.45, 1.55, 1.05, [[("Venue map", 13, True, False, WHITE)]],
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
arr = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.80), Inches(8.72), Inches(0.50), Inches(0.50))
arr.name = "VG arrow"
arr.fill.solid()
arr.fill.fore_color.rgb = WATCHC
arr.line.fill.background()
chip = card(s3, "VG chip", 6.40, 8.30, 1.55, 1.30, WHITE)
chip.line.color.rgb = CRIT
chip.line.width = Pt(1.5)
textbox(s3, "VG chipt", 6.45, 8.35, 1.45, 1.20, [
    [("Jam traced", 12, True, False, NAVY)],
    [("to Gate", 12, True, False, NAVY)],
], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

# ================= SLIDE 4: benchmark clips =================
s4 = p.slides[3]
kill(s4, "Crowd 106")
card(s4, "BC card", 10.20, 7.45, 9.00, 3.05, CARD)
textbox(s4, "BC title", 10.50, 7.58, 8.40, 0.45, [
    [("Tested on three real crowds", 16, True, False, NAVY)],
])
clips = [
    ("Station concourse", "16 people peak", "37 FPS", 10.50),
    ("Busy crossing", "7 people peak", "42 FPS", 13.35),
    ("Temple street", "15 people peak", "47 FPS", 16.20),
]
for (nm, pk, fps, x) in clips:
    card(s4, "BC chip " + nm, x, 8.10, 2.70, 1.55, WHITE)
    textbox(s4, "BC chipt " + nm, x + 0.05, 8.12, 2.60, 1.51, [
        [(nm, 13, True, False, NAVY)],
        [(pk, 11, False, False, GREY)],
        [(fps, 16, True, False, NAVY)],
    ], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
textbox(s4, "BC foot", 10.50, 9.72, 8.40, 0.62, [
    [("Same laptop, 500 frames per clip \u2014 model and tracker time only.", 10.5, False, False, GREY)],
])

p.save(SRC)
print("saved", SRC)
