"""Judge-hard pass 3: align deck copy with Problem #2 (density/movement/bottleneck/proactive).
Crowd-first headline, bottleneck steps, remove false 'all cameras' claim.
"""
import shutil
from pptx import Presentation

SRC = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
shutil.copy2(SRC, r"C:\Users\shoun\AppData\Local\Temp\opencode\simp-backup3.pptx")

SWAPS = {
    # slide 2: headline maps straight onto the problem statement
    "Know the crowd. Know the person. Stop the surge.":
        "Know the crowd. Catch the bottleneck. Stop the surge.",
    "Cameras watch. The software remembers faces.":
        "Cameras measure how dense each zone is, and how fast it moves.",
    # slide 2 watch-it-work: crowd first, face demoted to on-demand
    "A face appears on camera.":
        "Live feed: every person counted, zone density tracked.",
    "Their file pops up in under a second.":
        "A zone packs up and stops: bottleneck flagged.",
    "A crowded zone turns yellow in under 1 second.":
        "Guard warned in seconds; face file opens on demand.",
    # slide 3 plain words
    "In plain words: the camera watches, the software follows every person, and it warns guards before a crush builds.":
        "In plain words: the camera watches, the software flags where the crowd stops moving, and it warns guards before a crush builds.",
    # slide 4 red alert card: proactive countdown
    "Fires when a crowd packs tight and stops moving.":
        "Fires when a zone packs tight and stops. The wall counts down the seconds to critical.",
    # slide 5: remove false multi-camera claim
    "One screen watches all cameras and warns only about risky zones.":
        "One screen ranks every zone by risk, so guards only look where it matters.",
}

p = Presentation(SRC)
done, miss = set(), []
for slide in p.slides:
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            txt = "".join(r.text for r in para.runs).strip()
            if txt in SWAPS and txt not in done:
                repl = SWAPS[txt]
                if para.runs:
                    para.runs[0].text = repl
                    for r in para.runs[1:]:
                        r.text = ""
                    done.add(txt)
for k in SWAPS:
    if k not in done:
        miss.append(k)
p.save(SRC)
print(f"swapped {len(done)}/{len(SWAPS)}")
if miss:
    print("MISSING:")
    for m in miss:
        print(" -", m[:80])
