# Judge pass 2: fix overclaims, UI mismatches, PS-echo, citations
from pptx import Presentation

DECK = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
p = Presentation(DECK)
ok, fail = [], []


def swap(tf, old, new):
    for para in tf.paragraphs:
        t = "".join(r.text for r in para.runs)
        if t.strip() == old.strip():
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            return True
    return False


def edit_shape(slide, prefix, pairs):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.text_frame.text.startswith(prefix):
            for old, new in pairs:
                if swap(sh.text_frame, old, new):
                    ok.append(old[:50])
                else:
                    fail.append(prefix + " | " + old[:50])
            return
    fail.append("SHAPE NOT FOUND: " + prefix)


# --- Slide 2: legend must match live UI (Calm / Watch / Critical) ---
s2 = p.slides[1]
edit_shape(s2, "Busy", [("Busy", "Watch")])
edit_shape(s2, "Danger", [("Danger", "Critical")])

# --- Slide 3: cut "850 sightings" (ID-fragmentation landmine) ---
s3 = p.slides[2]
edit_shape(s3, "Tested on real video", [
    ("Followed 18 people at once and 850 sightings across 2 minutes of station video.",
     "Followed 18 people at once across 2 minutes of busy station video."),
])

# --- Slide 5 ---
s5 = p.slides[4]
edit_shape(s5, "Before", [
    ("50 screens, one tired guard, and risks get missed.",
     "A wall of screens, one tired guard, and risks get missed."),
])
edit_shape(s5, "After\nOne screen", [
    ("One screen ranks every zone by risk, so guards only look where it matters.",
     "One screen ranks every zone by risk: security moves from reactive to proactive."),
])
edit_shape(s5, "Police", [
    ("Find a lost child or a suspect from a photo in seconds. Full history in one click.",
     "A lost child's enrolled face is flagged live in seconds. Full history in one click."),
])
edit_shape(s5, "After\nA sealed", [
    ("After", "Handover"),
])
edit_shape(s5, "Pilot targets", [
    ("Pilot targets: photo search in seconds, guard warned in under 1 second, every step logged.",
     "Pilot targets: face matched live in seconds, guard warned in under 1 second, every step logged."),
])

# --- Slide 6: kill unsourced citation-flavored claims, name LFW, fix Hathras, fix title ---
s6 = p.slides[5]
edit_shape(s6, "RESEARCH", [
    ("RESEARCH\xa0 AND REFERENCES", "RESEARCH AND REFERENCES"),
])
edit_shape(s6, "Finding and following people", [
    ("Finds every person and follows each one. Uses YOLOv8 and ByteTrack, rated in 2025-26 studies as the best mix of speed and accuracy.",
     "Finds every person and follows each one. Uses YOLOv8 and ByteTrack, the standard real-time pairing for crowd tracking."),
])
edit_shape(s6, "Matching faces", [
    ("Compares a face with saved photos. Uses ArcFace and Faiss, which score 99.5 percent on the standard face test.",
     "Compares a face with saved photos. Uses ArcFace and Faiss: 99.5 percent accuracy on LFW, the standard face test."),
])
edit_shape(s6, "Why India needs it", [
    ("Hathras, Kumbh and railway station crushes all came from packed halls and blocked exits.",
     "Packed crowds and blocked exits turned deadly at Hathras, Kumbh and railway stations."),
])

p.save(DECK)
print("OK:", len(ok))
for x in ok:
    print("  +", x)
print("FAIL:", len(fail))
for x in fail:
    print("  -", x)
