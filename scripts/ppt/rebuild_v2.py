"""rebuild_v2: full content rebuild within template boundaries.
Keeps: slide 1, watermark groups, title boxes (text swapped), footers, card containers.
Replaces: all body content with crowd-physics narrative. No laptop scores. No banned cites."""
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
shutil.copy(SRC, r"C:\Users\shoun\AppData\Local\Temp\opencode\simp-backup8.pptx")

NAVY = RGBColor(0x1E, 0x27, 0x61)
BODY = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xEE, 0xF2, 0xFF)
GOLD = RGBColor(0xE5, 0xB8, 0x22)

p = Presentation(SRC)
LINK_COUNT = [0]


def kill(slide, names):
    for nm in names:
        found = False
        for sh in list(slide.shapes):
            if sh.name == nm:
                sh._element.getparent().remove(sh._element)
                found = True
        print(("killed " if found else "MISSING ") + nm)


def kill_prefix(slide, prefix):
    n = 0
    for sh in list(slide.shapes):
        if sh.name.startswith(prefix):
            sh._element.getparent().remove(sh._element)
            n += 1
    print(f"killed {n} with prefix {prefix}")


def run_fmt(r, txt, sz, bold, italic, color, url=None):
    r.text = txt
    f = r.font
    f.name = "Calibri"
    f.size = Pt(sz)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    if url:
        try:
            r.hyperlink.address = url
            LINK_COUNT[0] += 1
        except Exception as e:
            print("link fail", url, e)


def set_box(shape, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT):
    """paras: list of paragraphs; each a list of (txt, sz, bold, italic, color, url?)."""
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for pi, runs in enumerate(paras):
        para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        para.alignment = align
        for run in runs:
            url = run[5] if len(run) > 5 else None
            r = para.add_run()
            run_fmt(r, run[0], run[1], run[2], run[3], run[4], url)


def card_text(shape, title, tsz, body, bsz, links=None, tcolor=NAVY, bcolor=BODY):
    paras = [[(title, tsz, True, False, tcolor)]]
    if body:
        paras.append([(body, bsz, False, False, bcolor)])
    if links:
        paras.append([(d + ("  " if i < len(links) - 1 else ""), 11, False, False, GREY, u)
                       for i, (d, u) in enumerate(links)])
    set_box(shape, paras)


def retitle(slide, new):
    for sh in slide.shapes:
        if sh.name == "TextBox 4" and sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                if para.text.strip():
                    para.runs[0].text = new
                    for r in para.runs[1:]:
                        r.text = ""
                    print("retitled to", new)
                    return
    print("TITLE NOT FOUND")


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    print("MISSING", name)
    return None


def textbox(slide, name, l, t, w, h, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT):
    s = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    s.name = name
    set_box(s, paras, anchor, align)
    return s


def panel(slide, name, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    try:
        s.adjustments[0] = 0.07
    except Exception:
        pass
    s.text_frame.clear()
    return s


# ================= SLIDE 2: THE PROBLEM =================
s2 = p.slides[1]
retitle(s2, "THE PROBLEM")
set_box(by_name(s2, "TextBox 101"), [
    [("Guards watch screens. Danger builds unseen.", 32, True, False, NAVY)],
    [("A control room watches dozens of flat camera feeds. A guard spots trouble only after "
      "the crowd stops moving. By then, crushing pressure is already passing body to body. "
      "Stillness is not the warning. It is the accident.", 20, False, False, BODY)],
])
card_text(by_name(s2, "Stat 102"), "Kumbh 2025", 20,
          "2,751 cameras and AI alerts were live. The crush still killed 30.", 15)
card_text(by_name(s2, "Stat 103"), "Itaewon 2022", 20,
          "159 dead in a 3.2 metre alley. 11 calls ignored. No machine spoke up.", 15)
card_text(by_name(s2, "Stat 104"), "7 per m2", 20,
          "A crowd turns fluid. Shockwaves throw people metres. (Fruin)", 15)
kill(s2, ["TextBox 105", "Num 106", "TextBox 107", "Num 108", "TextBox 109",
          "Num 110", "TextBox 111", "Legend 114", "TextBox 115", "Dot 116",
          "TextBox 117", "Dot 118", "TextBox 119", "Dot 120", "TextBox 121"])
kill_prefix(s2, "VM ")
panel(s2, "P2 time", 9.80, 2.85, 9.40, 5.00, CARD)
textbox(s2, "P2 timet", 10.10, 3.00, 8.80, 0.60,
        [[("Itaewon, 29 October 2022", 20, True, False, NAVY)]])
events = [
    ("18:34", "First distress calls from the packed alley.", 3.80),
    ("hours pass", "Calls keep coming. No machine raises an alarm.", 5.00),
    ("22:15", "The alley locks. 159 dead, 196 injured.", 6.20),
]
for (tm, ds, y) in events:
    d = s2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.30), Inches(y + 0.12), Inches(0.55), Inches(0.55))
    d.name = "P2 dot"
    d.fill.solid()
    d.fill.fore_color.rgb = NAVY
    d.line.fill.background()
    textbox(s2, "P2 ev", 11.05, y, 7.60, 0.95, [
        [(tm, 16, True, False, NAVY)],
        [(ds, 14, False, False, BODY)],
    ])
set_box(by_name(s2, "TextBox 113"),
        [[("Four hours of human warning. No machine ever paged anyone.", 16, False, True, GREY)]])
panel(s2, "P2 claim", 9.80, 8.55, 9.40, 1.90, NAVY)
set_box(by_name(s2, "P2 claim"), [
    [("Danger is invisible pressure, not visible stillness.", 18, True, False, WHITE)],
    [("Measure pressure, warn early, open gates.", 15, False, False, WHITE)],
])

# ================= SLIDE 3: HOW IT WORKS =================
s3 = p.slides[2]
retitle(s3, "HOW IT WORKS")
steps = [
    ("Chevron 101", "Calibrate", "TextBox 102", "pixels to ground, once per camera"),
    ("Chevron 103", "Map density", "TextBox 104", "heatmaps count overlap"),
    ("Chevron 105", "Track flow", "TextBox 106", "flow reads crowd currents"),
    ("Chevron 107", "Measure pressure", "TextBox 108", "density times movement"),
    ("Chevron 109", "Forecast jam", "TextBox 110", "inflow beats outflow"),
    ("Chevron 111", "Warn guards", "TextBox 112", "open gates upstream"),
]
for (cn, ct, tn, tt) in steps:
    c = by_name(s3, cn)
    if c:
        set_box(c, [[(ct, 20, True, False, WHITE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    t = by_name(s3, tn)
    if t:
        set_box(t, [[(tt, 15, False, False, GREY)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
set_box(by_name(s3, "Banner 113"), [[(
    "In plain words: map the ground right, count the crowd as heat, read its "
    "currents, warn on pressure.", 20, False, False, NAVY)]])
card_text(by_name(s3, "Card 115"), "Heatmaps beat boxes", 18,
          "At crush density a detector scores F1 0.07 against 0.71 point-based. (NWPU-Crowd)", 13.5)
card_text(by_name(s3, "Card 116"), "Pressure, not stillness", 18,
          "Pressure is density times uneven movement. Density or speed alone cannot warn. (Johansson 2008)", 13.5)
card_text(by_name(s3, "Card 117"), "Side track, kept", 18,
          "Below 2 people per m2, box IDs still run and feed the dossier. Already built.", 13.5)
kill_prefix(s3, "VG ")
panel(s3, "P3 mae", 0.80, 6.55, 7.40, 4.00, CARD)
textbox(s3, "P3 maet", 1.05, 6.70, 6.90, 0.55,
        [[("Count error keeps falling (MAE, lower is better)", 15, True, False, NAVY)]])
bars = [("CSRNet", 68, 7.45), ("Bayesian", 63, 8.25), ("DM-Count", 60, 9.05), ("Latest", 49, 9.85)]
for (lbl, val, y) in bars:
    textbox(s3, "P3 lbl", 1.05, y, 1.70, 0.55, [[(lbl, 13, False, False, NAVY)]],
            anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
    tr = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.85), Inches(y + 0.10), Inches(3.60), Inches(0.32))
    tr.name = "P3 track"
    tr.fill.solid()
    tr.fill.fore_color.rgb = RGBColor(0xD5, 0xDB, 0xEA)
    tr.line.fill.background()
    fl = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.85), Inches(y + 0.10),
                             Inches(3.60 * val / 70.0), Inches(0.32))
    fl.name = "P3 fill"
    fl.fill.solid()
    fl.fill.fore_color.rgb = NAVY
    fl.line.fill.background()
    textbox(s3, "P3 val", 6.60, y, 0.90, 0.55, [[(str(val), 13, True, False, NAVY)]],
            anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT)

# ================= SLIDE 4: FEASIBILITY =================
s4 = p.slides[3]
card_text(by_name(s4, "Card 103"), "Video stays put", 20,
          "Each edge box sends small JSON numbers, never raw video. Footage is discarded at the camera.", 15)
card_text(by_name(s4, "Card 104"), "Calibrated where it counts", 20,
          "Fixed gates are measured once. Other views use relative levels. No fake square metres.", 15)
card_text(by_name(s4, "Card 105"), "Private by design", 20,
          "Heatmaps cannot show a face. Saved video is blurred, wiped after 7 days, zero identity in core.", 15)
kill(s4, ["Chart 4", "TextBox 102"])
kill_prefix(s4, "BC ")
cc = by_name(s4, "ChartCard 101")
textbox(s4, "P4 edget", 1.10, 3.05, 8.40, 0.60,
        [[("One camera, one edge box", 20, True, False, NAVY)]])
nodes = [
    ("Camera: fixed gate view", NAVY, 3.80),
    ("Edge box: density plus flow plus pressure on device", NAVY, 4.95),
    ("Sends JSON numbers only, never video", GOLD, 6.10),
    ("Venue server: fuses every camera into one map", NAVY, 7.25),
    ("Guard wall: ranked zones, early warning", NAVY, 8.40),
]
for (txt, fill, y) in nodes:
    nd = panel(s4, "P4 node", 1.55, y, 7.50, 0.80, fill)
    tc = WHITE
    set_box(nd, [[(txt, 15, True, False, tc)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    if y < 8.40:
        ar = s4.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(4.80), Inches(y + 0.82), Inches(0.40), Inches(0.28))
        ar.name = "P4 arrow"
        ar.fill.solid()
        ar.fill.fore_color.rgb = NAVY
        ar.line.fill.background()
textbox(s4, "P4 edgef", 1.10, 9.45, 8.40, 0.80,
        [[("Raw video never leaves the camera box.", 14, False, False, GREY)]])
panel(s4, "P4 site", 10.20, 7.45, 9.00, 3.05, CARD)
textbox(s4, "P4 sitet", 10.50, 7.58, 8.40, 0.45,
        [[("Runs on site, not in the cloud", 16, True, False, NAVY)]])
srows = [
    "Edge box per 1 to 4 cameras (Jetson class)",
    "Venue server fuses every feed into one map",
    "Guard handsets get the alert, human confirms",
]
for i, r in enumerate(srows):
    textbox(s4, "P4 siterow", 10.50, 8.12 + i * 0.62, 8.40, 0.60,
            [[(r, 13.5, False, False, BODY)]])
textbox(s4, "P4 sitef", 10.50, 9.98, 8.40, 0.40,
        [[("Deployment scale numbers only, no laptop tests.", 10.5, False, False, GREY)]])

# ================= SLIDE 5: IMPACT =================
s5 = p.slides[4]
set_box(by_name(s5, "TextBox 107"), [[("When pressure rises", 20, True, False, NAVY)]])
set_box(by_name(s5, "TextBox 109"), [
    [("Yellow watch", 20, True, False, WHITE)],
    [("Guards watch the zone closely. Overflow gates get ready.", 15, False, False, WHITE)],
])
set_box(by_name(s5, "TextBox 111"), [
    [("Red alert", 20, True, False, WHITE)],
    [("Guards get steps on their handsets; a human confirms before any PA or door release.", 15, False, False, WHITE)],
])
set_box(by_name(s5, "TextBox 113"), [
    [("Handover", 20, True, False, WHITE)],
    [("A sealed evidence pack goes to the police.", 15, False, False, WHITE)],
])
set_box(by_name(s5, "TextBox 115"), [[(
    "Pilot targets: warn before critical, page guards in seconds, log every step.",
    15, True, False, NAVY)]])
set_box(by_name(s5, "After 103"), [
    [("After", 18, True, False, GOLD)],
    [("Pressure rising at Gate B opens overflow at Gate A upstream. Security acts before the jam, not after.", 16, True, False, WHITE)],
])

# ================= SLIDE 6: REFERENCES =================
s6 = p.slides[5]
set_box(by_name(s6, "TextBox 101"), [[("The science behind it", 20, True, False, NAVY)]])
card_text(by_name(s6, "Card 102"), "Count the overlap", 18,
          "CSRNet (CVPR 2018), Bayesian Loss (ICCV 2019), DM-Count (NeurIPS 2020).",
          13, links=[
              ("arXiv:1802.10062", "https://arxiv.org/abs/1802.10062"),
              ("arXiv:1908.03684", "https://arxiv.org/abs/1908.03684"),
              ("DM-Count", "https://github.com/cvlab-stonybrook/DM-Count")])
card_text(by_name(s6, "Card 103"), "Map pixels to ground", 18,
          "Chan perspective maps (CVPR 2008), Zhang ground-plane fusion (CVPR 2019).",
          13, links=[
              ("people counting paper", "http://www.svcl.ucsd.edu/publications/conference/2008/cvpr08/cvpr08_peoplecnt.pdf"),
              ("ground-plane paper", "https://openaccess.thecvf.com/content_CVPR_2019/html/Zhang_Wide-Area_Crowd_Counting_via_Ground-Plane_Density_Maps_and_Multi-View_Fusion_CVPR_2019_paper.html")])
card_text(by_name(s6, "Card 104"), "Read the currents", 18,
          "Mehran social force (CVPR 2009), streaklines (ECCV 2010).",
          13, links=[
              ("CVPR 2009 pdf", "https://vision.eecs.ucf.edu/papers/cvpr2009/CVPR09_Mehran.pdf"),
              ("streakline project", "https://www.crcv.ucf.edu/projects/streakline_eccv")])
card_text(by_name(s6, "Card 105"), "Warn on pressure", 18,
          "Helbing social force (1995), crowd disasters (2007), Johansson pressure (2008).",
          13, links=[
              ("DOI:10.1103/PhysRevE.51.4282", "https://doi.org/10.1103/PhysRevE.51.4282"),
              ("arXiv:0708.3339", "https://arxiv.org/abs/0708.3339"),
              ("pressure analysis", "https://ar5iv.labs.arxiv.org/html/0810.4590")])
set_box(by_name(s6, "TextBox 106"), [[("Disasters, limits, footage", 20, True, False, NAVY)]])
card_text(by_name(s6, "Card 107"), "Itaewon and Love Parade", 16,
          "159 dead, Seoul 2022, ignored calls. 21 dead, Duisburg 2010. Turbulence, not panic.", 13)
card_text(by_name(s6, "Card 108"), "Kumbh and Hathras", 16,
          "30 dead, Jan 2025, with AI live. 121 dead, Jul 2024, 80k permit vs 250k crowd.", 13)
card_text(by_name(s6, "Card 109"), "Danger thresholds", 16,
          "Fluid at 7 per m2. Serious above 4. Flow peaks at 2 to 3. (Fruin, Still)", 13)
card_text(by_name(s6, "Card 110"), "Test footage", 16,
          "Public clips: station, crossing, temple.", 13, links=[
              ("Wikimedia Commons", "https://commons.wikimedia.org/")])

p.save(SRC)
print("saved", SRC, "links:", LINK_COUNT[0])
