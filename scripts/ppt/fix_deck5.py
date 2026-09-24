# Judge panel pass 3: strip face from core pitch, density+forecast MVP, honest scale/occlusion, Phase-2 face as B2G
from pptx import Presentation

DECK = r"D:\Q_project\_publish\deck\Raksha_Hackathon2026_Simplified.pptx"
p = Presentation(DECK)
ok, fail = [], []


def swap(tf, old, new):
    for para in tf.paragraphs:
        t = "".join(r.text for r in para.runs)
        if t.strip() == old.strip():
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            return True
    return False


def edit_shape(slide, prefix, pairs):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.text_frame.text.startswith(prefix):
            for old, new in pairs:
                if swap(sh.text_frame, old, new):
                    ok.append(old[:55])
                else:
                    fail.append(prefix + " | " + old[:55])
            return
    fail.append("SHAPE NOT FOUND: " + prefix)


# --- Slide 2: demo step 3 drops face, leads with forecast (proactive proof) ---
s2 = p.slides[1]
edit_shape(s2, "Guard warned in seconds", [
    ("Guard warned in seconds; face file opens on demand.",
     "Guard warned in seconds; wall shows time left to critical."),
])

# --- Slide 3: pipeline is density/flow/forecast, not faces ---
s3 = p.slides[2]
edit_shape(s3, "Know", [
    ("notes faces and clothes", "measures density and flow"),
])
edit_shape(s3, "Remember", [
    ("Remember", "Predict"),
    ("checks against saved photos", "forecasts time to critical"),
])
edit_shape(s3, "Keeps only good photos", [
    ("Keeps only good photos", "Crowd metrics, not faces"),
    ("Only sharp, front-facing photos are saved. Blurry ones are ignored, so matches stay reliable.",
     "Core computes density, flow and stillness per zone. No identity is needed for a crush warning."),
])
edit_shape(s3, "Works with normal cameras", [
    ("Any ordinary camera feed can be plugged in. No special equipment needed.",
     "Any ordinary camera works. High, wide angles give the clearest density per square metre."),
])

# --- Slide 4: face-speed card becomes forecast card; scale answered ---
s4 = p.slides[3]
edit_shape(s4, "Instant", [
    ("Instant", "Forecast"),
    ("A face is matched in just 0.06 seconds.",
     "The trend is fitted live, so the wall shows seconds to critical before the jam locks."),
])
edit_shape(s4, "Runs on one laptop", [
    ("Runs on one laptop", "Built to scale out"),
    ("Even the heaviest crowd runs smoothly.",
     "One laptop runs a venue today; an edge box per camera scales to mass gatherings."),
])

# --- Slide 5: police card loses face search; pilot targets forecast ---
s5 = p.slides[4]
edit_shape(s5, "Police", [
    ("A lost child's enrolled face is flagged live in seconds. Full history in one click.",
     "Every alert, zone and action exports as one evidence pack in a click."),
])
edit_shape(s5, "Pilot targets", [
    ("Pilot targets: face matched live in seconds, guard warned in under 1 second, every step logged.",
     "Pilot targets: tens of seconds forecast before critical, guard warned under 1 second, every step logged."),
])

# --- Slide 6: face card becomes Phase-2 B2G; privacy cards no longer contradict ---
s6 = p.slides[5]
edit_shape(s6, "Matching faces", [
    ("Matching faces", "Phase 2: face search"),
    ("Compares a face with saved photos. Uses ArcFace and Faiss: 99.5 percent accuracy on LFW, the standard face test.",
     "Not in the MVP. The pipeline is modular, so agencies can plug in their own face-search API later for post-incident cases (opt-in, B2G)."),
])
edit_shape(s6, "Blurred video, sealed proof", [
    ("Blurred video, sealed proof", "Anonymized by default"),
    ("Saved video is face-blurred. Sharp stills stay sealed in the case file.",
     "Saved video is face-blurred. Live alerts show zones and counts, not names."),
])
edit_shape(s6, "No ID linking", [
    ("No ID linking", "No ID in core"),
    ("Only people who choose to opt in can be matched.",
     "Crowd mode runs with zero identity. Face search stays opt-in Phase 2 only."),
])

p.save(DECK)
print("OK:", len(ok))
for x in ok:
    print("  +", x)
print("FAIL:", len(fail))
for x in fail:
    print("  -", x)
